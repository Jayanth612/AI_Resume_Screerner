from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


MD_PATH = Path(r"c:\Documents\IARE\project\Major Project\final\ResumeScreener\IEEE_PAPER_DRAFT.md")
DOCX_PATH = MD_PATH.with_name(f"{MD_PATH.stem}_IEEE.docx")


def set_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, start=60, bottom=60, end=60):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_paragraph(doc, text, style=None, align=None, bold=False, italic=False, size=12):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    return p


def parse_table(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", c) for c in cells):
            continue
        rows.append(cells)
    return rows


def main():
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    for style_name in ["Heading 1", "Heading 2"]:
        if style_name in doc.styles:
            doc.styles[style_name].font.name = "Times New Roman"
            doc.styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # Cover page
    title = lines[0].lstrip("# ").strip() if lines and lines[0].startswith("#") else "Resume Screening Paper"
    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_title.add_run(title)
    set_font(run, size=18, bold=True)

    cover_spacer = doc.add_paragraph()
    cover_spacer.paragraph_format.space_after = Pt(12)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("IEEE Conference Paper Draft")
    set_font(run, size=12, italic=True)

    doc.add_paragraph()
    doc.add_paragraph()

    meta_lines = [
        "Project: Resume Screener",
        "Prepared for: IEEE Conference Submission",
        "Author: [Add your name here]",
        "Institution: [Add your institution here]",
        "Date: 2026-07-28",
    ]
    for meta in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(meta)
        set_font(r, size=12)

    doc.add_page_break()

    i = 1
    in_equation = False

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("**Table"):
            add_paragraph(doc, stripped.strip("*"), align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
            i += 1
            continue

        if stripped == "$$":
            in_equation = not in_equation
            i += 1
            continue

        if in_equation:
            add_paragraph(doc, stripped, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_table(table_lines)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                table.autofit = True
                for r_idx, row in enumerate(rows):
                    for c_idx, cell_text in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = cell_text
                        set_cell_margins(cell)
                        if r_idx == 0:
                            shade_cell(cell, "D9E1F2")
                        for para in cell.paragraphs:
                            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in para.runs:
                                set_font(run, size=10.5, bold=(r_idx == 0))
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(re.sub(r"^\d+\.\s+", "", stripped))
            set_font(run)
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(stripped[2:].strip())
            set_font(run)
            i += 1
            continue

        add_paragraph(doc, stripped)
        i += 1

    for para in doc.paragraphs:
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.0

    doc.save(str(DOCX_PATH))
    print(f"Saved {DOCX_PATH}")


if __name__ == "__main__":
    main()